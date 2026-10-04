"""UI-independent MTAR workflow services.

Everything the Streamlit app does to an assessment goes through this module:
Gmail intake, lab import, photos, COC generation, and report generation. The
UI only collects input and displays results, so the workflow rules stay in one
place and can be tested without a browser.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import threading
import time
from uuid import uuid4


from coc_builder import build_coc_payload, fill_coc_pdf, validate_coc_payload
from document_store import FileDocumentStore
from drive_photos import (
    DriveClient,
    folder_id_from_link,
    folder_link,
    folder_match_score,
    match_job_folders,
    suggested_folders,
    sync_job_photos,
)
from gmail_intake import (
    GmailApiClient,
    confident_job_match,
    find_job_by_report_number,
    inspect_attachment,
    is_confident_prolab_result,
)
from job_store import SQLiteJobStore, new_persistent_job
from models import new_area, new_sample
from pdf_report_builder import create_customer_pdf
from photo_service import normalize_report_photo
from prolab_parser import (
    _parse_lab_date,
    apply_prolab_results,
    build_automated_job_from_prolab,
    parse_prolab_pdf,
    suggested_mapping,
)
from report_builder import MOLD_DESCRIPTIONS, create_report
from workflow import complete_job, final_report_issues, reopen_job, transition_job


RAILWAY_VOLUME = os.environ.get("RAILWAY_VOLUME_MOUNT_PATH")
DEFAULT_DATA_ROOT = Path(RAILWAY_VOLUME) if RAILWAY_VOLUME else Path(".")
DB_PATH = Path(os.environ.get("MTAR_DB_PATH", str(DEFAULT_DATA_ROOT / "mtar_jobs.sqlite3")))
DOCUMENT_ROOT = Path(os.environ.get("MTAR_DOCUMENT_ROOT", str(DEFAULT_DATA_ROOT / "mtar_data")))
COC_TEMPLATE = Path(os.environ.get("MTAR_COC_TEMPLATE", "assets/BLANK_COC.pdf"))
UPLOADED_COC_TEMPLATE = DOCUMENT_ROOT / "templates" / "BLANK_COC.pdf"
GMAIL_POLL_SECONDS = max(60, int(os.environ.get("GMAIL_POLL_SECONDS", "300")))

store = SQLiteJobStore(DB_PATH)
documents = FileDocumentStore(DOCUMENT_ROOT)
gmail_client = GmailApiClient()
drive_client = DriveClient()
DRIVE_PHOTOS_FOLDER = folder_id_from_link(os.environ.get("DRIVE_PHOTOS_FOLDER", ""))

# Manual "Check Gmail Now" and the background poller must never import the
# same message twice by running at the same moment.
INTAKE_LOCK = threading.Lock()
LAST_GMAIL_CHECK: dict = {"status": "not_run"}

STATUS_LABELS = {
    "draft": "Draft",
    "inspection": "Inspection",
    "awaiting_lab": "Awaiting Lab",
    "lab_received": "Lab Received",
    "report_review": "Report Review",
    "ready_to_send": "Ready to Send",
    "sent": "Sent",
    "closed": "Completed",
}

FINDING_OPTIONS = [
    "Needs consultant review",
    "Active mold growth confirmed",
    "Elevated spore counts",
    "Mold levels not elevated",
    "Visual mold present",
    "No mold detected",
]

REPORT_OUTCOMES = [
    "Pending consultant review",
    "Mold remediation required",
    "No significant mold contamination identified",
]

PHOTO_KINDS = {
    "sampling": "Sampling",
    "inspection": "Moisture assessment",
    "thermal": "Thermal Imaging",
}

# The two moisture-assessment sentences the inspector picks between per area.
MOISTURE_STATES = {
    "dry": "No surfaces were wet in this area.",
    "wet": "Surfaces were wet in this area.",
}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def status_label(status: str) -> str:
    return STATUS_LABELS.get(status, str(status or "").replace("_", " ").title())


def safe_filename(value: str, fallback: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in "._-" else "_" for c in (value or "")).strip("._")
    return cleaned or fallback


def client_file_stem(job: dict) -> str:
    return safe_filename(job.get("client_name", ""), "Client")


def sample_name(sample: dict) -> str:
    if sample.get("outdoor_control"):
        return sample.get("name") or "Outdoor Control"
    return sample.get("name") or sample.get("location") or "Unnamed Sample"


def save_job(job: dict) -> dict:
    return store.save(job)


def moisture_sentence(state: str | None, highest: float | None, extra: str = "") -> str:
    """Report text for an area's moisture assessment, e.g.
    "No surfaces were wet in this area. Highest moisture content observed was 14%."
    """
    parts = []
    if state in MOISTURE_STATES:
        parts.append(MOISTURE_STATES[state])
        if highest is not None:
            parts.append(f"Highest moisture content observed was {float(highest):g}%.")
    if str(extra or "").strip():
        parts.append(str(extra).strip())
    return " ".join(parts)


def advance_status_if(job: dict, current_status: str, target_status: str) -> dict:
    if job.get("status") != current_status:
        return job
    try:
        return transition_job(job, target_status)
    except ValueError:
        return job


# ---------------------------------------------------------------------------
# Assessment creation
# ---------------------------------------------------------------------------


def create_field_job() -> dict:
    """Create a clean inspection-first assessment with only the outdoor control."""
    job = new_persistent_job()
    job["areas"] = []
    outdoor = new_sample(
        sample_type="Air Sample",
        location="Outdoor Control",
        outdoor_control=True,
    )
    outdoor.update(
        {
            "name": "Outdoor Control",
            "serial_number": "",
            "sample_type_code": "P15",
            "flow_rate_liters": 15,
            "flow_rate_minutes": 5,
            "mold_analysis": True,
        }
    )
    job["samples"] = [outdoor]
    job["air_lab_rows"] = []
    job["surface_lab_rows"] = []
    job["mold_types"] = []
    return job


def _link_profiles(job: dict) -> None:
    """Create or reuse the customer and property profiles for an assessment."""
    try:
        customer = store.ensure_customer(job.get("client_name", ""))
        property_record = store.ensure_property(
            customer["id"],
            job.get("address", ""),
            job.get("city", ""),
            job.get("state", ""),
            job.get("zip", ""),
        )
        job["customer_id"] = customer["id"]
        job["property_id"] = property_record["id"]
    except ValueError:
        # A valid report is still retained for review if a future laboratory
        # format omits a customer or address.
        pass


def _generate_auto_drafts(job: dict, lab_pdf: bytes | None) -> None:
    """Generate the AUTO_DRAFT DOCX and PDF right after a lab import.

    Lab-only drafts contain explicit review-required placeholders for
    inspection facts that are not available in the PRO-LAB PDF.
    """
    try:
        report = create_report(job, report_photos(job), lab_pdf)
        draft_filename = f"{client_file_stem(job)}_Mold_Assessment_AUTO_DRAFT.docx"
        documents.save_bytes(job["id"], "reports", draft_filename, report.getvalue())
        job["latest_draft_filename"] = draft_filename

        pdf = create_customer_pdf(job, report_photos(job))
        pdf_filename = f"{client_file_stem(job)}_Mold_Assessment_AUTO_DRAFT.pdf"
        documents.save_bytes(job["id"], "reports", pdf_filename, pdf.getvalue())
        job["latest_pdf_draft_filename"] = pdf_filename
        job.pop("automatic_draft_error", None)
    except Exception as exc:
        job["automatic_draft_error"] = str(exc)


def import_lab_report(
    pdf_bytes: bytes,
    filename: str,
    parsed: dict,
    *,
    all_jobs: list[dict],
    gmail_source: dict | None = None,
) -> tuple[dict, bool]:
    """Attach a confident PRO-LAB report to an assessment, creating one if needed.

    High-confidence matches attach to an existing Awaiting Lab assessment.
    Otherwise MTAR creates a new assessment directly from the lab report, so
    the lab report can be the trigger for the workflow. Returns the saved
    assessment and whether it was newly created. Callers check report-number
    duplicates first.
    """
    awaiting_jobs = [job for job in all_jobs if job.get("status") == "awaiting_lab"]
    match = confident_job_match(parsed, awaiting_jobs)
    target = None
    created = False
    match_score = 0
    match_reasons = ["created from PRO-LAB report"]

    if match:
        target = next((job for job in awaiting_jobs if job.get("id") == match["job_id"]), None)
        if target:
            match_score = match["score"]
            match_reasons = match["reasons"]

    if target is None:
        target = new_persistent_job()
        mapping = build_automated_job_from_prolab(target, parsed, MOLD_DESCRIPTIONS.keys())
        target["lab_mapping"] = mapping
        target["status"] = "report_review"
        created = True
    else:
        mapping = suggested_mapping(parsed, target)
        target["lab_mapping"] = mapping
        parsed_keys = {sample.get("key") for sample in parsed.get("samples", [])}
        mapped_keys = {key for key, value in mapping.items() if value}
        if parsed_keys and parsed_keys == mapped_keys and len(set(mapping.values())) == len(mapping):
            apply_prolab_results(target, parsed, mapping, MOLD_DESCRIPTIONS.keys())
            target["status"] = "report_review"
        else:
            try:
                target = transition_job(target, "lab_received")
            except ValueError:
                target["status"] = "lab_received"

    _link_profiles(target)

    filename = safe_filename(filename, "PROLAB_Result.pdf")
    documents.save_bytes(target["id"], "lab", filename, pdf_bytes)
    target["lab_filename"] = filename
    target["lab_parsed"] = parsed
    if gmail_source:
        target.setdefault("gmail_message_ids", []).append(gmail_source["message_id"])
        target["gmail_message_ids"] = list(dict.fromkeys(target["gmail_message_ids"]))
        target["gmail_source"] = {
            **gmail_source,
            "filename": filename,
            "match_score": match_score,
            "match_reasons": match_reasons,
            "created_assessment": created,
        }

    _generate_auto_drafts(target, pdf_bytes)
    return store.save(target), created


def _all_jobs() -> list[dict]:
    return [job for summary in store.list(limit=500) if (job := store.get(summary.id))]


def _is_partial_prolab(parsed: dict) -> bool:
    """A PDF that has some PRO-LAB structure but failed the confidence check."""
    metadata = parsed.get("metadata", {})
    return bool(parsed.get("samples") or metadata.get("report_number"))


# New PRO-LAB reports wait on the PRO-LAB Reports page until the consultant
# creates an assessment, unless this setting is switched on.
AUTO_CREATE_SETTING = "auto_create_from_lab"
AUTO_CREATE_DEFAULT = False


def auto_create_enabled() -> bool:
    return store.get_setting(AUTO_CREATE_SETTING, "1" if AUTO_CREATE_DEFAULT else "0") == "1"


def set_auto_create(enabled: bool) -> None:
    store.set_setting(AUTO_CREATE_SETTING, "1" if enabled else "0")


def _lab_details(metadata: dict, samples: int | None = None) -> dict:
    details = {
        "test_location": str(metadata.get("test_location") or ""),
        "report_date": str(metadata.get("report_date") or ""),
    }
    if samples is not None:
        details["samples"] = samples
    return details


def run_gmail_intake() -> dict:
    """Check Gmail once and turn valid PRO-LAB PDFs into MTAR assessments.

    A report for an assessment that is Awaiting Lab attaches to it. Other
    reports wait on the PRO-LAB Reports page, unless auto-create is on.
    Uncertain PRO-LAB-like PDFs are listed there as needing review, and
    unrelated PDFs are ignored.
    """
    if not gmail_client.configured():
        return {
            "configured": False,
            "checked": 0,
            "imported": 0,
            "created": 0,
            "ambiguous": [],
            "ignored": 0,
            "skipped": 0,
            "waiting": 0,
        }

    with INTAKE_LOCK:
        all_jobs = _all_jobs()
        processed_ids = {
            message_id
            for job in all_jobs
            for message_id in job.get("gmail_message_ids", [])
        }
        processed_ids |= store.inbox_message_ids()
        deleted_reports = store.deleted_report_numbers()

        result = {
            "configured": True,
            "checked": 0,
            "imported": 0,
            "created": 0,
            "ambiguous": [],
            "ignored": 0,
            "skipped": 0,
            "waiting": 0,
        }
        auto_create = auto_create_enabled()

        attachments = gmail_client.search_pdf_attachments()
        result["checked"] = len(attachments)

        for attachment in attachments:
            if attachment.message_id in processed_ids:
                result["skipped"] += 1
                continue

            inspected = inspect_attachment(attachment)
            parsed = inspected["parsed"]
            metadata = parsed.get("metadata", {})

            if not parsed.get("samples") or not is_confident_prolab_result(parsed):
                filename = safe_filename(attachment.filename, "Lab_Result.pdf")
                if _is_partial_prolab(parsed):
                    # Keep the PDF so the consultant can match or create from it.
                    documents.save_bytes(f"inbox_{attachment.message_id}", "lab", filename, attachment.content)
                    store.record_inbox_item(
                        attachment.message_id,
                        "needs_review",
                        filename=filename,
                        sender=attachment.sender,
                        subject=attachment.subject,
                        report_number=str(metadata.get("report_number") or ""),
                        project_name=str(metadata.get("project_name") or ""),
                        reason="; ".join(parsed.get("warnings", []))
                        or "Report number and sample COC lines do not agree.",
                        details=_lab_details(metadata, len(parsed.get("samples", []))),
                    )
                    result["ambiguous"].append(
                        {
                            "subject": attachment.subject,
                            "project_name": metadata.get("project_name", ""),
                            "report_number": metadata.get("report_number", ""),
                        }
                    )
                else:
                    store.record_inbox_item(
                        attachment.message_id,
                        "ignored",
                        filename=filename,
                        sender=attachment.sender,
                        subject=attachment.subject,
                        reason="Not a PRO-LAB result",
                    )
                    result["ignored"] += 1
                processed_ids.add(attachment.message_id)
                continue

            # An assessment the consultant deleted stays deleted, even when
            # the same lab report arrives again in another email.
            report_key = " ".join(str(metadata.get("report_number") or "").upper().split())
            if report_key and report_key in deleted_reports:
                store.record_inbox_item(
                    attachment.message_id,
                    "deleted",
                    filename=safe_filename(attachment.filename, "Lab_Result.pdf"),
                    sender=attachment.sender,
                    subject=attachment.subject,
                    report_number=report_key,
                    reason="Assessment for this lab report was deleted",
                )
                processed_ids.add(attachment.message_id)
                result["skipped"] += 1
                continue

            # Duplicate protection: the same PRO-LAB report number never
            # creates a second assessment, even from a different email.
            existing = find_job_by_report_number(parsed, all_jobs)
            if existing is None:
                report_number = str(metadata.get("report_number") or "").strip()
                existing = next(
                    (
                        job for job in all_jobs
                        if report_number
                        and str(job.get("lab_metadata", {}).get("report_number") or "").strip() == report_number
                    ),
                    None,
                )
            if existing:
                existing.setdefault("gmail_message_ids", []).append(attachment.message_id)
                existing["gmail_message_ids"] = list(dict.fromkeys(existing["gmail_message_ids"]))
                store.save(existing)
                processed_ids.add(attachment.message_id)
                result["skipped"] += 1
                continue

            # Without auto-create, a report only attaches by itself to an
            # assessment that is waiting for it; anything else waits for the
            # consultant on the PRO-LAB Reports page.
            awaiting_jobs = [job for job in all_jobs if job.get("status") == "awaiting_lab"]
            if not auto_create and not confident_job_match(parsed, awaiting_jobs):
                filename = safe_filename(attachment.filename, "PROLAB_Result.pdf")
                documents.save_bytes(f"inbox_{attachment.message_id}", "lab", filename, attachment.content)
                store.record_inbox_item(
                    attachment.message_id,
                    "waiting",
                    filename=filename,
                    sender=attachment.sender,
                    subject=attachment.subject,
                    report_number=str(metadata.get("report_number") or ""),
                    project_name=str(metadata.get("project_name") or ""),
                    details=_lab_details(metadata, len(parsed.get("samples", []))),
                )
                processed_ids.add(attachment.message_id)
                result["waiting"] += 1
                continue

            saved, created = import_lab_report(
                attachment.content,
                attachment.filename,
                parsed,
                all_jobs=all_jobs,
                gmail_source={
                    "message_id": attachment.message_id,
                    "thread_id": attachment.thread_id,
                    "sender": attachment.sender,
                    "subject": attachment.subject,
                },
            )
            all_jobs = [job for job in all_jobs if job.get("id") != saved["id"]] + [saved]
            processed_ids.add(attachment.message_id)
            result["imported"] += 1
            if created:
                result["created"] += 1

        return result


def check_gmail_now() -> dict:
    """Run one intake pass and record the outcome for the dashboard."""
    global LAST_GMAIL_CHECK
    checked_at = datetime.now(timezone.utc)
    try:
        result = run_gmail_intake()
    except Exception as exc:
        LAST_GMAIL_CHECK = {"status": "error", "error": str(exc), "at": checked_at}
        raise
    if result.get("configured"):
        LAST_GMAIL_CHECK = {
            "status": "ok",
            "at": checked_at,
            "checked": result.get("checked", 0),
            "imported": result.get("imported", 0),
            "created": result.get("created", 0),
            "ambiguous": len(result.get("ambiguous", [])),
            "ignored": result.get("ignored", 0),
            "waiting": result.get("waiting", 0),
        }
    return result


def gmail_poll_loop() -> None:
    """Background polling loop. Streamlit starts it once per server process."""
    while True:
        if gmail_client.configured():
            try:
                check_gmail_now()
            except Exception:
                pass  # recorded in LAST_GMAIL_CHECK
        if drive_client.configured():
            try:
                sync_all_drive_photos()
            except Exception:
                pass  # per-assessment errors are recorded on the assessment
        time.sleep(GMAIL_POLL_SECONDS)


def create_from_lab_upload(pdf_bytes: bytes, filename: str) -> dict:
    """Manual equivalent of the Gmail trigger: upload a PRO-LAB PDF.

    Returns {"status": "created" | "attached" | "duplicate" | "invalid", ...}.
    """
    parsed = parse_prolab_pdf(pdf_bytes)
    if not parsed.get("samples") or not is_confident_prolab_result(parsed):
        reason = "; ".join(parsed.get("warnings", [])) or (
            "This PDF does not look like a complete PRO-LAB result (report number and sample table must agree)."
        )
        return {"status": "invalid", "reason": reason}

    with INTAKE_LOCK:
        all_jobs = _all_jobs()
        existing = find_job_by_report_number(parsed, all_jobs)
        if existing:
            return {"status": "duplicate", "job_id": existing["id"]}
        saved, created = import_lab_report(pdf_bytes, filename, parsed, all_jobs=all_jobs)
    return {"status": "created" if created else "attached", "job_id": saved["id"]}


# ---------------------------------------------------------------------------
# Lab Inbox (uncertain PRO-LAB-like PDFs)
# ---------------------------------------------------------------------------


def inbox_pdf(item: dict) -> bytes | None:
    return documents.read_bytes(f"inbox_{item['message_id']}", "lab", item.get("filename", ""))


def ignore_inbox_item(message_id: str) -> None:
    store.resolve_inbox_item(message_id, "dismissed")


def attach_inbox_item_to_job(message_id: str, job_id: str) -> dict:
    """Consultant decided which assessment an uncertain lab PDF belongs to."""
    item = store.get_inbox_item(message_id)
    job = store.get(job_id)
    if not item or not job:
        raise ValueError("Lab report or assessment not found")
    content = inbox_pdf(item)
    if content is None:
        raise ValueError("The stored lab PDF is missing")
    job = attach_lab_pdf(job, content, item.get("filename") or "PROLAB_Result.pdf")
    job.setdefault("gmail_message_ids", []).append(message_id)
    job["gmail_message_ids"] = list(dict.fromkeys(job["gmail_message_ids"]))
    job = store.save(job)
    store.resolve_inbox_item(message_id, "matched", job_id=job["id"])
    return job


def create_job_from_inbox_item(message_id: str) -> dict:
    """Consultant confirmed an uncertain PDF should start a new assessment."""
    item = store.get_inbox_item(message_id)
    if not item:
        raise ValueError("Lab report not found")
    content = inbox_pdf(item)
    if content is None:
        raise ValueError("The stored lab PDF is missing")
    parsed = parse_prolab_pdf(content)
    if not parsed.get("samples"):
        raise ValueError("No sample table could be read from this PDF")
    existing = find_job_by_report_number(parsed, _all_jobs())
    if existing:
        raise ValueError(
            f"PRO-LAB report #{parsed['metadata']['report_number']} already has an assessment: "
            f"{existing.get('client_name') or 'New Assessment'}."
        )
    from_gmail = not message_id.startswith("deleted:")
    with INTAKE_LOCK:
        saved, _ = import_lab_report(
            content,
            item.get("filename") or "PROLAB_Result.pdf",
            parsed,
            all_jobs=[],  # explicit create: never auto-attach elsewhere
            gmail_source={
                "message_id": message_id,
                "thread_id": "",
                "sender": item.get("sender", ""),
                "subject": item.get("subject", ""),
            } if from_gmail else None,
        )
    store.resolve_inbox_item(message_id, "created", job_id=saved["id"])
    return try_auto_link_drive_folder(saved)


# ---------------------------------------------------------------------------
# Dashboard housekeeping: complete, reopen, delete, finished reports
# ---------------------------------------------------------------------------

FINISHED_REPORT_TYPES = (".pdf", ".docx")


def lab_report_number(job: dict) -> str:
    metadata = job.get("lab_metadata") or (job.get("lab_parsed") or {}).get("metadata") or {}
    return str(metadata.get("report_number") or "").strip()


def complete_assessment(job: dict) -> dict:
    """Mark an assessment completed, whatever step it is on."""
    return store.save(complete_job(job))


def reopen_assessment(job: dict) -> dict:
    return store.save(reopen_job(job))


def delete_assessment(job_id: str) -> bool:
    """Delete an assessment with all its files, for good.

    Its Gmail messages and lab report number are remembered, so the Gmail
    check does not create the same assessment again.
    """
    with INTAKE_LOCK:
        job = store.get(job_id)
        if not job:
            return False
        metadata = job.get("lab_metadata") or (job.get("lab_parsed") or {}).get("metadata") or {}
        lab_filename = safe_filename(job.get("lab_filename") or "", "PROLAB_Result.pdf")
        details = {
            "report_number": lab_report_number(job),
            "project_name": str(metadata.get("project_name") or job.get("client_name", "")),
            "subject": ", ".join(p for p in [job.get("address"), job.get("city")] if p),
            "reason": "Assessment deleted in MTAR",
            "filename": lab_filename,
            "details": _lab_details(metadata),
        }
        for message_id in job.get("gmail_message_ids", []):
            store.record_inbox_item(message_id, "deleted", **details)
            store.resolve_inbox_item(message_id, "deleted")
        # Keep the lab PDF so the assessment can be created again from the
        # PRO-LAB Reports page.
        lab_pdf = lab_pdf_bytes(job)
        if lab_pdf:
            documents.save_bytes(f"inbox_deleted:{job_id}", "lab", lab_filename, lab_pdf)
        store.record_inbox_item(f"deleted:{job_id}", "deleted", **details)
        store.delete(job_id)
    documents.delete_job_files(job_id)
    return True


def attach_finished_report(job: dict, content: bytes, filename: str, *, complete: bool = True) -> dict:
    """Store the consultant's own finished report on an assessment."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in FINISHED_REPORT_TYPES:
        raise ValueError("Upload the finished report as a PDF or Word (.docx) file.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    name = safe_filename(Path(filename).name, f"{client_file_stem(job)}_Finished_Report{suffix}")
    documents.save_bytes(job["id"], "finished", name, content)
    job["finished_report_filename"] = name
    job["finished_report_uploaded_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    if complete:
        job = complete_job(job)
    return store.save(job)


def _finished_report_text(content: bytes, filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    text = ""
    try:
        if suffix == ".pdf":
            import fitz

            with fitz.open(stream=content, filetype="pdf") as doc:
                text = " ".join(page.get_text() for page in list(doc)[:3])
        elif suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(content))
            parts = [p.text for p in document.paragraphs[:120]]
            for table in document.tables[:6]:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
            text = " ".join(parts)
    except Exception:
        text = ""
    # The file name often carries the customer's name, so it counts too.
    return f"{Path(filename or '').stem.replace('_', ' ')} {text}"


def _match_key(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() else " " for ch in str(value or "").upper())
    return f" {' '.join(cleaned.split())} "


def match_finished_report(content: bytes, filename: str, jobs: list[dict]) -> tuple[dict | None, str]:
    """Find the one assessment a finished report belongs to.

    A report matches when it contains the assessment's house number and
    street name. When several assessments share that address, the customer's
    name must also appear. Returns (assessment or None, reason).
    """
    text = _match_key(_finished_report_text(content, filename))
    if " CERTIFICATE OF MOLD ANALYSIS " in text and " MOLD ASSESSMENT REPORT " not in text:
        return None, "this is a PRO-LAB lab report, not a finished assessment report"

    by_address = []
    for job in jobs:
        words = _match_key(job.get("address")).split()
        if len(words) < 2 or not words[0][:1].isdigit():
            continue
        if f" {words[0]} {words[1]} " in text:
            by_address.append(job)
    if not by_address:
        return None, "no assessment address was found in it"
    if len(by_address) == 1:
        return by_address[0], ""

    by_name = [
        job for job in by_address
        if job.get("client_name") and _match_key(job["client_name"]) in text
    ]
    if len(by_name) == 1:
        return by_name[0], ""
    open_jobs = [job for job in (by_name or by_address) if job.get("status") != "closed"]
    if len(open_jobs) == 1:
        return open_jobs[0], ""
    return None, f"{len(by_address)} assessments have this address"


def import_finished_reports(files: list[tuple[str, bytes]]) -> list[dict]:
    """Attach several finished reports to their assessments and complete them."""
    jobs = _all_jobs()
    results = []
    for filename, content in files:
        job, reason = match_finished_report(content, filename, jobs)
        if job is None:
            results.append({"filename": filename, "job_id": None, "reason": reason})
            continue
        try:
            saved = attach_finished_report(job, content, filename)
        except ValueError as exc:
            results.append({"filename": filename, "job_id": None, "reason": str(exc)})
            continue
        jobs = [saved if j["id"] == saved["id"] else j for j in jobs]
        results.append({"filename": filename, "job_id": saved["id"], "client_name": saved.get("client_name", ""), "reason": ""})
    return results


# ---------------------------------------------------------------------------
# PRO-LAB Reports page: every lab report MTAR has seen
# ---------------------------------------------------------------------------

LAB_REPORT_STATUSES = {
    "new": "New",
    "review": "Needs review",
    "linked": "Has assessment",
    "deleted": "Assessment deleted",
    "dismissed": "Dismissed",
}
_INBOX_TO_REPORT_STATUS = {
    "waiting": "new",
    "needs_review": "review",
    "created": "linked",
    "matched": "linked",
    "deleted": "deleted",
    "dismissed": "dismissed",
}
_REPORT_PRIORITY = ["linked", "new", "review", "deleted", "dismissed"]


def lab_reports() -> list[dict]:
    """One entry per PRO-LAB report: Gmail ones, uploads, and deleted ones.

    Waiting and uncertain Gmail reports come from the Lab Inbox table; reports
    already turned into assessments come from the assessments themselves, so
    the list also covers assessments created before this page existed.
    """
    jobs = {job["id"]: job for job in _all_jobs()}
    entries: list[dict] = []
    for item in store.list_inbox(status=list(_INBOX_TO_REPORT_STATUS)):
        status = _INBOX_TO_REPORT_STATUS[item["status"]]
        job_id = item.get("job_id") or ""
        if status == "linked" and job_id not in jobs:
            continue  # the assessment itself is listed below, or was deleted
        if not documents.exists(f"inbox_{item['message_id']}", "lab", item.get("filename") or ""):
            continue  # nothing to show or create from
        try:
            details = json.loads(item.get("details") or "{}")
        except ValueError:
            details = {}
        entries.append({
            "key": item["message_id"],
            "message_id": item["message_id"],
            "job_id": job_id if status == "linked" else "",
            "status": status,
            "customer": item.get("project_name") or "",
            "location": details.get("test_location", ""),
            "report_number": item.get("report_number") or "",
            "report_date": _parse_lab_date(details.get("report_date", "")),
            "filename": item.get("filename") or "",
            "reason": item.get("reason") or "",
            "found_at": item.get("created_at") or "",
        })
    for job in jobs.values():
        if not job.get("lab_filename") or not lab_pdf_present(job):
            continue
        metadata = job.get("lab_metadata") or (job.get("lab_parsed") or {}).get("metadata") or {}
        entries.append({
            "key": f"job:{job['id']}",
            "message_id": None,
            "job_id": job["id"],
            "status": "linked",
            "customer": str(metadata.get("project_name") or job.get("client_name") or ""),
            "location": str(metadata.get("test_location") or ", ".join(p for p in [job.get("address"), job.get("city")] if p)),
            "report_number": lab_report_number(job),
            "report_date": _parse_lab_date(str(metadata.get("report_date") or "")),
            "filename": job["lab_filename"],
            "reason": "",
            "found_at": job.get("created_at") or "",
        })

    best: dict[str, dict] = {}
    for entry in entries:
        key = " ".join(entry["report_number"].upper().split()) or entry["key"]
        current = best.get(key)
        if current is None or _REPORT_PRIORITY.index(entry["status"]) < _REPORT_PRIORITY.index(current["status"]):
            best[key] = entry
    return sorted(best.values(), key=lambda e: (e["report_date"] or datetime.min.date(), e["found_at"]), reverse=True)


def lab_report_pdf(entry: dict) -> bytes | None:
    if entry.get("message_id"):
        return documents.read_bytes(f"inbox_{entry['message_id']}", "lab", entry.get("filename") or "")
    job = store.get(entry.get("job_id") or "")
    return lab_pdf_bytes(job) if job else None


def create_assessment_from_report(entry: dict) -> dict:
    """The consultant chose to start an assessment from this lab report."""
    if entry.get("status") == "linked":
        raise ValueError("This lab report already has an assessment.")
    return create_job_from_inbox_item(entry["message_id"])


# ---------------------------------------------------------------------------
# Areas and samples
# ---------------------------------------------------------------------------


def add_area(job: dict, name: str | None = None) -> dict:
    job.setdefault("areas", []).append(new_area(name or f"Inspection Area {len(job.get('areas', [])) + 1}"))
    return store.save(job)


def remove_area(job: dict, area_id: str) -> dict:
    job["areas"] = [area for area in job.get("areas", []) if area.get("id") != area_id]
    for sample in job.get("samples", []):
        if sample.get("area_id") == area_id:
            sample["area_id"] = None
    retained = []
    for photo in job.get("photos", []):
        if photo.get("area_id") == area_id:
            documents.delete(job["id"], "photos", photo.get("filename", ""))
        else:
            retained.append(photo)
    job["photos"] = retained
    return store.save(job)


def add_sample(job: dict, sample_type: str) -> dict:
    area_id = job["areas"][0].get("id") if job.get("areas") else None
    sample = new_sample(sample_type=sample_type, area_id=area_id)
    if sample_type == "Air Sample":
        sample.update({"sample_type_code": "P15", "flow_rate_liters": 15, "flow_rate_minutes": 5})
    else:
        sample["sample_type_code"] = "SW"
    sample.update({"serial_number": "", "mold_analysis": True})
    job.setdefault("samples", []).append(sample)
    return store.save(job)


def remove_sample(job: dict, sample_id: str) -> dict:
    if sample_id == "sample_outdoor_control":
        raise ValueError("The outdoor control sample is kept as a protected baseline.")
    job["samples"] = [s for s in job.get("samples", []) if s.get("id") != sample_id]
    job["air_lab_rows"] = [r for r in job.get("air_lab_rows", []) if r.get("sample_id") != sample_id]
    job["surface_lab_rows"] = [r for r in job.get("surface_lab_rows", []) if r.get("sample_id") != sample_id]
    return store.save(job)


# ---------------------------------------------------------------------------
# Photos
# ---------------------------------------------------------------------------


def add_photo(
    job: dict,
    raw: bytes,
    source_name: str,
    role: str,
    *,
    area_id: str | None = None,
    kind: str = "inspection",
    drive_file_id: str | None = None,
    drive_path: str = "",
) -> dict:
    """Normalize and store one photo. A new property photo replaces the old one."""
    normalized = normalize_report_photo(raw)
    photos = job.setdefault("photos", [])
    if role == "property":
        for existing in [p for p in photos if p.get("role") == "property"]:
            documents.delete(job["id"], "photos", existing.get("filename", ""))
        job["photos"] = photos = [p for p in photos if p.get("role") != "property"]

    token = uuid4().hex[:8]
    stem = Path(safe_filename(source_name, f"photo_{len(photos) + 1}.jpg")).stem
    prefix = role if role in {"property", "outdoor", "environment", "unsorted"} else f"area_{area_id}"
    filename = f"{prefix}_{token}_{stem}.jpg"
    documents.save_bytes(job["id"], "photos", filename, normalized)
    photos.append(
        {
            "id": f"photo_{token}",
            "filename": filename,
            "role": role,
            "area_id": area_id,
            "caption": "",
            "kind": kind,
            **({"drive_file_id": drive_file_id, "drive_path": drive_path} if drive_file_id else {}),
        }
    )
    return store.save(job)


def delete_photo(job: dict, photo_id: str) -> dict:
    target = next((p for p in job.get("photos", []) if p.get("id") == photo_id), None)
    if target:
        documents.delete(job["id"], "photos", target.get("filename", ""))
        job["photos"] = [p for p in job.get("photos", []) if p.get("id") != photo_id]
    return store.save(job)


def assign_photo(job: dict, photo_id: str, role: str, area_id: str | None = None) -> dict:
    """File an unsorted photo under a report section."""
    if role == "area" and not any(a.get("id") == area_id for a in job.get("areas", [])):
        raise ValueError("Choose an inspection area for this photo.")
    if role not in {"property", "outdoor", "environment", "area"}:
        raise ValueError("Choose where this photo goes.")
    target = next((p for p in job.get("photos", []) if p.get("id") == photo_id), None)
    if not target:
        raise ValueError("Photo not found.")
    if role == "property":  # only one cover photo is kept
        for existing in [p for p in job["photos"] if p.get("role") == "property" and p is not target]:
            documents.delete(job["id"], "photos", existing.get("filename", ""))
        job["photos"] = [p for p in job["photos"] if p.get("role") != "property" or p is target]
    target["role"] = role
    target["area_id"] = area_id if role == "area" else None
    target["kind"] = {"outdoor": "sampling", "environment": "environment"}.get(role, "inspection")
    return store.save(job)


def photo_bytes(job: dict, photo: dict) -> bytes | None:
    return documents.read_bytes(job["id"], "photos", photo.get("filename", ""))


def report_photos(job: dict) -> dict:
    """Group stored photos by the slot the report generator places them in."""
    result: dict = {}
    for photo in job.get("photos", []):
        content = photo_bytes(job, photo)
        if not content:
            continue
        entry = {
            "content": BytesIO(content),
            "caption": photo.get("caption", ""),
            "kind": photo.get("kind", "inspection"),
        }
        role = photo.get("role")
        if role in {"property", "outdoor", "environment"}:
            result.setdefault(role, []).append(entry)
        elif role == "area" and photo.get("area_id"):
            result.setdefault(photo["area_id"], []).append(entry)
    return result


# ---------------------------------------------------------------------------
# Lab results on an existing assessment
# ---------------------------------------------------------------------------


def attach_lab_pdf(job: dict, pdf_bytes: bytes, filename: str) -> dict:
    """Store a lab PDF on this assessment and suggest sample matches."""
    parsed = parse_prolab_pdf(pdf_bytes)
    if not parsed.get("samples"):
        raise ValueError("; ".join(parsed.get("warnings", [])) or "No structured PRO-LAB result table was found.")
    filename = safe_filename(filename, "PROLAB_Result.pdf")
    documents.save_bytes(job["id"], "lab", filename, pdf_bytes)
    job["lab_filename"] = filename
    job["lab_parsed"] = parsed
    job["lab_mapping"] = suggested_mapping(parsed, job)
    job = advance_status_if(job, "awaiting_lab", "lab_received")
    return store.save(job)


def mapping_issues(job: dict) -> list[str]:
    parsed_samples = (job.get("lab_parsed") or {}).get("samples", [])
    mapping = job.get("lab_mapping") or {}
    issues = []
    if any(not mapping.get(lab.get("key")) for lab in parsed_samples):
        issues.append("Map every lab sample before applying results.")
    mapped = [mapping[lab["key"]] for lab in parsed_samples if mapping.get(lab.get("key"))]
    if len(mapped) != len(set(mapped)):
        issues.append("Each lab sample must map to a different MTAR sample.")
    return issues


def apply_lab(job: dict) -> dict:
    issues = mapping_issues(job)
    if issues:
        raise ValueError(" ".join(issues))
    apply_prolab_results(job, job.get("lab_parsed") or {}, job.get("lab_mapping") or {}, MOLD_DESCRIPTIONS.keys())
    job = advance_status_if(job, "lab_received", "report_review")
    return store.save(job)


def lab_pdf_bytes(job: dict) -> bytes | None:
    filename = job.get("lab_filename")
    return documents.read_bytes(job["id"], "lab", filename) if filename else None


def lab_pdf_present(job: dict) -> bool:
    return bool(job.get("lab_filename")) and documents.exists(job["id"], "lab", job.get("lab_filename", ""))


# ---------------------------------------------------------------------------
# COC
# ---------------------------------------------------------------------------


def coc_template_path() -> Path | None:
    for path in (COC_TEMPLATE, UPLOADED_COC_TEMPLATE):
        if path.exists():
            return path
    return None


def save_coc_template(content: bytes) -> None:
    UPLOADED_COC_TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    UPLOADED_COC_TEMPLATE.write_bytes(content)


def generate_coc(job: dict) -> tuple[dict, str, bytes]:
    payload = build_coc_payload(job)
    issues = validate_coc_payload(payload)
    if issues:
        raise ValueError("COC needs: " + "; ".join(issues))
    template = coc_template_path()
    if not template:
        raise ValueError("Upload the blank PRO-LAB COC template before generating a COC.")
    output = fill_coc_pdf(template.read_bytes(), payload)
    filename = f"{client_file_stem(job)}_PROLAB_COC.pdf"
    documents.save_bytes(job["id"], "coc", filename, output)
    job["latest_coc_filename"] = filename
    return store.save(job), filename, output


# ---------------------------------------------------------------------------
# Reports
# ---------------------------------------------------------------------------


def final_issues(job: dict) -> list[str]:
    return final_report_issues(job, lab_pdf_present=lab_pdf_present(job))


def generate_review_docx(job: dict) -> tuple[dict, str, bytes]:
    output = create_report(job, report_photos(job), lab_pdf_bytes(job)).getvalue()
    filename = f"{client_file_stem(job)}_Mold_Assessment_DRAFT.docx"
    documents.save_bytes(job["id"], "reports", filename, output)
    job["latest_draft_filename"] = filename
    return store.save(job), filename, output


def generate_review_pdf(job: dict) -> tuple[dict, str, bytes]:
    output = create_customer_pdf(job, report_photos(job)).getvalue()
    filename = f"{client_file_stem(job)}_Mold_Assessment_DRAFT.pdf"
    documents.save_bytes(job["id"], "reports", filename, output)
    job["latest_pdf_draft_filename"] = filename
    return store.save(job), filename, output


def generate_final_docx(job: dict) -> tuple[dict, str, bytes]:
    issues = final_issues(job)
    if issues:
        raise ValueError("Final report blocked: " + "; ".join(issues))
    output = create_report(job, report_photos(job), lab_pdf_bytes(job)).getvalue()
    filename = f"{client_file_stem(job)}_Mold_Assessment_FINAL.docx"
    documents.save_bytes(job["id"], "reports", filename, output)
    job["latest_final_filename"] = filename
    job = advance_status_if(job, "report_review", "ready_to_send")
    return store.save(job), filename, output


def generate_final_pdf(job: dict) -> tuple[dict, str, bytes]:
    """Customer PDF. The PRO-LAB certificate is emailed as its own attachment."""
    issues = final_issues(job)
    if issues:
        raise ValueError("Final PDF blocked: " + "; ".join(issues))
    output = create_customer_pdf(job, report_photos(job)).getvalue()
    filename = f"{client_file_stem(job)}_Mold_Assessment_FINAL.pdf"
    documents.save_bytes(job["id"], "reports", filename, output)
    job["latest_final_pdf_filename"] = filename
    job = advance_status_if(job, "report_review", "ready_to_send")
    return store.save(job), filename, output


# ---------------------------------------------------------------------------
# Google Drive photos
# ---------------------------------------------------------------------------

DRIVE_SYNC_STATUSES = {"draft", "inspection", "awaiting_lab", "lab_received", "report_review"}


def link_drive_folder(job: dict, link: str) -> dict:
    folder_id = folder_id_from_link(link)
    if not folder_id:
        raise ValueError("Paste a Google Drive folder link (it contains /folders/).")
    job["drive_folder_id"] = folder_id
    job.pop("drive_folder_name", None)
    return store.save(job)


def use_drive_folder(job: dict, folder: dict) -> dict:
    """Link a folder picked from the Jobs folder list."""
    job["drive_folder_id"] = folder["id"]
    job["drive_folder_name"] = folder.get("name", "")
    return store.save(job)


def unlink_drive_folder(job: dict) -> dict:
    job.pop("drive_folder_id", None)
    job.pop("drive_folder_name", None)
    return store.save(job)


_DRIVE_ROOT_CACHE: dict = {"at": 0.0, "items": []}


def drive_job_folders(max_age: float = 60) -> list[dict]:
    """The folders inside the Drive Jobs folder (re-read at most once a minute)."""
    if not (DRIVE_PHOTOS_FOLDER and drive_client.configured()):
        return []
    if time.monotonic() - _DRIVE_ROOT_CACHE["at"] > max_age:
        _DRIVE_ROOT_CACHE["items"] = drive_client.list_children(DRIVE_PHOTOS_FOLDER)
        _DRIVE_ROOT_CACHE["at"] = time.monotonic()
    return list(_DRIVE_ROOT_CACHE["items"])


def drive_folder_choices(job: dict) -> list[dict]:
    """Job folders for the picker, the likeliest first, without folders other assessments use."""
    taken = {other.get("drive_folder_id") for other in _all_jobs() if other.get("id") != job.get("id")}
    return [folder for folder in suggested_folders(job, drive_job_folders()) if folder["id"] not in taken]


def auto_link_drive_folders(root_children: list[dict] | None = None) -> list[dict]:
    """Link every open assessment that has one clear job folder in Drive."""
    if root_children is None:
        root_children = drive_job_folders(max_age=0)
    if not root_children:
        return []
    jobs = _all_jobs()
    taken = {job["drive_folder_id"] for job in jobs if job.get("drive_folder_id")}
    open_jobs = [job for job in jobs if job.get("status") in DRIVE_SYNC_STATUSES and not job.get("drive_folder_id")]
    linked = []
    for job_id, folder in match_job_folders(open_jobs, root_children, taken).items():
        job = store.get(job_id)
        if job and not job.get("drive_folder_id"):
            job["drive_folder_id"] = folder["id"]
            job["drive_folder_name"] = folder.get("name", "")
            linked.append(store.save(job))
    return linked


def try_auto_link_drive_folder(job: dict) -> dict:
    """Link a just-created assessment right away when its folder is clear."""
    try:
        with INTAKE_LOCK:
            auto_link_drive_folders()
    except Exception:
        return job  # the background check tries again
    return store.get(job["id"]) or job


def sync_drive_photos(job: dict, *, root_children: list[dict] | None = None) -> tuple[dict, dict]:
    """Import new photos from this assessment's Drive folder."""
    if not drive_client.configured():
        raise ValueError("Google Drive is not connected. Add the Drive scope to the Google sign-in first.")
    if not job.get("drive_folder_id"):
        auto_link_drive_folders(root_children)
        job = store.get(job["id"]) or job
    if not job.get("drive_folder_id"):
        raise ValueError("No Drive folder is linked to this assessment yet.")

    job, result = sync_job_photos(job, drive_client, add_photo=add_photo, add_area=add_area)
    summary = {
        "at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "imported": result.imported,
        "skipped": result.skipped,
        "areas_created": result.areas_created,
        "unplaced": result.unplaced,
        "errors": result.errors,
    }
    job["drive_last_sync"] = summary
    return store.save(job), summary


def sync_all_drive_photos() -> int:
    """Background pass: pull new Drive photos for every open assessment."""
    root_children = drive_job_folders(max_age=0) if DRIVE_PHOTOS_FOLDER else None
    if root_children:
        with INTAKE_LOCK:
            auto_link_drive_folders(root_children)
    imported = 0
    for summary in store.list(statuses=sorted(DRIVE_SYNC_STATUSES), limit=500):
        # The lock keeps a delete from the dashboard from running mid-sync,
        # which would otherwise save the deleted assessment back.
        with INTAKE_LOCK:
            job = store.get(summary.id)
            if not job or job.get("status") not in DRIVE_SYNC_STATUSES:
                continue
            if not job.get("drive_folder_id"):
                continue
            try:
                _, result = sync_drive_photos(job, root_children=root_children)
                imported += result["imported"]
            except Exception as exc:
                job["drive_last_sync"] = {"at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "errors": [str(exc)]}
                store.save(job)
    return imported

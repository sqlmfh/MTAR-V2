from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
import unittest
import tempfile

from PIL import Image

from job_store import SQLiteJobStore, new_persistent_job
from prolab_parser import build_automated_job_from_prolab
from report_builder import MOLD_DESCRIPTIONS, create_report
from pdf_report_builder import create_customer_pdf
import fitz


FIXTURE = Path(__file__).parent / "fixtures" / "scarlet_parsed_expected.json"


class EmailToAssessmentRegressionTests(unittest.TestCase):
    def setUp(self):
        self.parsed = json.loads(FIXTURE.read_text(encoding="utf-8"))

    def test_scarlet_lab_creates_assessment_areas_and_assigns_samples(self):
        job = new_persistent_job()
        mapping = build_automated_job_from_prolab(
            job,
            self.parsed,
            MOLD_DESCRIPTIONS.keys(),
        )

        self.assertEqual(job["client_name"], "Scarlet Harper")
        self.assertEqual(job["address"], "16371 County Road 245")
        self.assertEqual(job["city"], "Terrell")
        self.assertEqual(job["state"], "TX")
        self.assertEqual(job["zip"], "75160")
        self.assertEqual(job["lab_metadata"]["report_number"], "2030805")
        self.assertEqual(set(mapping), {"2030805-1", "2030805-2", "2030805-3"})

        self.assertEqual([area["name"] for area in job["areas"]], ["Coat Closet", "Bedroom Closet"])
        areas = {area["name"]: area for area in job["areas"]}
        self.assertEqual(set(areas), {"Bedroom Closet", "Coat Closet"})
        self.assertEqual(areas["Bedroom Closet"]["finding"], "Mold levels not elevated")
        self.assertEqual(areas["Coat Closet"]["finding"], "Active mold growth confirmed")
        self.assertEqual(areas["Bedroom Closet"]["description"], "")
        self.assertEqual(areas["Coat Closet"]["description"], "")
        self.assertIn("Air Sample:", areas["Bedroom Closet"]["lab_summary"])
        self.assertIn("Swab Sample:", areas["Coat Closet"]["lab_summary"])

        samples = {sample.get("lab_coc_line"): sample for sample in job["samples"]}
        outdoor = samples["2030805-1"]
        bedroom = samples["2030805-2"]
        coat = samples["2030805-3"]

        self.assertTrue(outdoor["outdoor_control"])
        self.assertEqual(outdoor["serial_number"], "Q2986725")
        self.assertEqual(outdoor["flow_rate_liters"], 15)
        self.assertEqual(outdoor["flow_rate_minutes"], 5)

        self.assertEqual(bedroom["serial_number"], "Q2986713")
        self.assertEqual(bedroom["area_id"], areas["Bedroom Closet"]["id"])
        self.assertEqual(bedroom["flow_rate_liters"], 15)
        self.assertEqual(bedroom["flow_rate_minutes"], 5)
        self.assertEqual(bedroom["lab_determination"], "NOT ELEVATED")

        self.assertEqual(coat["serial_number"], "CC")
        self.assertEqual(coat["area_id"], areas["Coat Closet"]["id"])
        self.assertEqual(coat["lab_determination"], "UNUSUAL")
        self.assertIn("Hyphae", coat["lab_fungi"])

        self.assertEqual(
            job["mold_types"],
            [
                "Cladosporium",
                "Curvularia",
                "Epicoccum",
                "Hyphae",
                "Other Ascospores",
                "Other Basidiospores",
                "Penicillium/Aspergillus",
                "Smuts, myxomycetes",
            ],
        )
        self.assertNotIn("Alternaria", job["mold_types"])

    def test_customer_and_property_profiles_are_reused(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = SQLiteJobStore(Path(tmp) / "mtar.sqlite3")
            customer_a = store.ensure_customer("Scarlet Harper")
            customer_b = store.ensure_customer("  SCARLET   HARPER  ")
            self.assertEqual(customer_a["id"], customer_b["id"])

            property_a = store.ensure_property(
                customer_a["id"],
                "16371 County Road 245",
                "Terrell",
                "TX",
                "75160",
            )
            property_b = store.ensure_property(
                customer_b["id"],
                "16371 COUNTY ROAD 245",
                "TERRELL",
                "tx",
                "75160",
            )
            self.assertEqual(property_a["id"], property_b["id"])

    def test_lab_only_assessment_can_generate_customer_pdf(self):
        job = new_persistent_job()
        build_automated_job_from_prolab(job, self.parsed, MOLD_DESCRIPTIONS.keys())

        pdf = create_customer_pdf(job, {})
        output = pdf.getvalue()

        self.assertTrue(output.startswith(b"%PDF"))
        doc = fitz.open(stream=output, filetype="pdf")
        try:
            self.assertGreaterEqual(len(doc), 8)
            first_page = doc[0].get_text("text")
            self.assertIn("MOLD ASSESSMENT REPORT", first_page.upper())
            self.assertIn("Scarlet Harper", first_page)
        finally:
            doc.close()


    def test_scarlet_photo_layout_generates_twelve_page_customer_report(self):
        job = new_persistent_job()
        build_automated_job_from_prolab(job, self.parsed, MOLD_DESCRIPTIONS.keys())
        job["report_outcome"] = "Mold remediation required"
        job["humidity"] = 44

        areas = {area["name"]: area for area in job["areas"]}
        areas["Coat Closet"]["moisture_notes"] = (
            "All surfaces were dry during the time of inspection. "
            "All materials were below 14% moisture content."
        )
        areas["Bedroom Closet"]["moisture_notes"] = (
            "All surfaces were dry during the time of inspection. "
            "All materials were below 14% moisture content."
        )
        areas["Coat Closet"]["thermal_notes"] = (
            "Thermal Imaging: No abnormalities were observed. "
            "All temperature variations were consistent with normal conditions."
        )
        areas["Bedroom Closet"]["thermal_notes"] = areas["Coat Closet"]["thermal_notes"]

        def photo():
            buffer = BytesIO()
            Image.new("RGB", (640, 480), "white").save(buffer, format="JPEG")
            buffer.seek(0)
            return {"content": buffer}

        def batch(count, kind="inspection"):
            rows = []
            for _ in range(count):
                item = photo()
                item["kind"] = kind
                rows.append(item)
            return rows

        photos = {
            "property": [photo()],
            "outdoor": batch(3, "sampling"),
            "environment": [photo()],
            areas["Coat Closet"]["id"]: (
                batch(3, "sampling") + batch(18, "inspection") + batch(4, "thermal")
            ),
            areas["Bedroom Closet"]["id"]: (
                batch(2, "sampling") + batch(6, "inspection") + batch(4, "thermal")
            ),
        }

        pdf = create_customer_pdf(job, photos)
        doc = fitz.open(stream=pdf.getvalue(), filetype="pdf")
        try:
            self.assertEqual(len(doc), 12)
            letter = doc[2].get_text("text")
            self.assertIn("Coat Closet", letter)
            self.assertNotIn("Bedroom Closet", letter)  # not elevated, so not listed

            outdoor = doc[3].get_text("text")
            self.assertLess(outdoor.index("Outdoor Control Sample"), outdoor.index("Visual Observations"))
            self.assertNotIn("Moisture Readings", outdoor)
            self.assertIn("COC / LINE #:", outdoor)

            coat = doc[4].get_text("text")
            self.assertIn("COAT CLOSET", coat.upper())
            self.assertLess(coat.index("Swab Sample:"), coat.index("Moisture Assessment:"))
            self.assertIn("2030805 - 3", coat)
            self.assertIn("Thermal Imaging", doc[5].get_text("text"))
            self.assertIn("BEDROOM CLOSET", doc[7].get_text("text").upper())
            self.assertIn("LABORATORY RESULTS ANALYSIS", doc[9].get_text("text").upper())
            self.assertIn("CONCLUSIONS", doc[10].get_text("text").upper())
            terms = doc[11].get_text("text")
            self.assertIn("TERMS AND CONDITIONS", terms.upper())
            self.assertNotIn("Laboratory Report Attached", terms)

            # Every photo is a 1.96in square, three to a row.
            coat_photos = [info["bbox"] for info in doc[4].get_image_info() if info["bbox"][0] < 540]
            for x0, y0, x1, y1 in coat_photos:
                self.assertAlmostEqual(x1 - x0, 1.96 * 72, delta=0.5)
                self.assertAlmostEqual(y1 - y0, 1.96 * 72, delta=0.5)
            first_row = [box for box in coat_photos if abs(box[1] - coat_photos[0][1]) < 1]
            self.assertEqual(len(first_row), 3)
        finally:
            doc.close()

    def test_word_report_follows_the_customer_report_layout(self):
        from docx import Document
        from docx.oxml.ns import qn

        job = new_persistent_job()
        build_automated_job_from_prolab(job, self.parsed, MOLD_DESCRIPTIONS.keys())
        job["report_outcome"] = "Mold remediation required"
        job["humidity"] = 58.0
        areas = {area["name"]: area for area in job["areas"]}
        for area in areas.values():
            area["moisture_notes"] = "No surfaces were wet in this area. Highest moisture content observed was 14%."

        def photo(kind):
            buffer = BytesIO()
            Image.new("RGB", (480, 640), "white").save(buffer, format="JPEG")
            buffer.seek(0)
            return {"content": buffer, "kind": kind}

        coat = areas["Coat Closet"]["id"]
        photos = {
            "outdoor": [photo("sampling") for _ in range(3)],
            "environment": [photo("environment")],
            coat: [photo("inspection") for _ in range(4)] + [photo("sampling") for _ in range(2)],
        }
        doc = Document(BytesIO(create_report(job, photos, None).getvalue()))
        texts = [p.text for p in doc.paragraphs]
        full = "\n".join(texts)

        self.assertNotIn("assigned to", full)
        self.assertNotIn("Laboratory Report Attached", full)
        self.assertNotIn("Moisture Readings", full)
        # The letter lists only areas where mold was found.
        bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet" and "—" in p.text]
        self.assertEqual(bullets, ["Coat Closet — Active mold growth confirmed"])

        def at(text, start=0):
            return next(i for i, t in enumerate(texts) if i >= start and t.startswith(text))

        outdoor = at("Outdoor Control Sample")
        self.assertLess(outdoor, at("Visual Observations"))
        self.assertEqual(texts[at("COC / LINE #:", outdoor)], "COC / LINE #: 2030805 - 1")

        # Coat Closet: bold "Swab Sample:", the sample's COC line and sampling
        # photos, then Moisture Assessment with the moisture photos.
        title = texts.index("Coat Closet")
        swab = at("Swab Sample:", title)
        self.assertTrue(doc.paragraphs[swab].runs[0].bold)
        self.assertEqual(texts[at("COC / LINE #:", title)], "COC / LINE #: 2030805 - 3")
        moisture = at("Moisture Assessment:", title)

        def photo_count(start, end):
            return sum(len(p._p.findall(".//" + qn("wp:inline"))) for p in doc.paragraphs[start:end])

        self.assertEqual(photo_count(swab, moisture), 2)
        self.assertEqual(photo_count(moisture, texts.index("Bedroom Closet")), 4)
        self.assertEqual(photo_count(outdoor, at("Visual Observations")), 3)
        for shape in doc.inline_shapes:
            if abs(shape.width.inches - 1.96) < 0.01:
                self.assertAlmostEqual(shape.height.inches, 1.96, places=2)

    def test_lab_only_assessment_can_generate_review_draft(self):
        job = new_persistent_job()
        build_automated_job_from_prolab(job, self.parsed, MOLD_DESCRIPTIONS.keys())

        report = create_report(job, {}, None)
        output = report.getvalue()

        self.assertGreater(len(output), 10000)
        self.assertTrue(output.startswith(b"PK"))
        self.assertEqual(job["report_outcome"], "Pending consultant review")


if __name__ == "__main__":
    unittest.main()

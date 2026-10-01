from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document

from models import new_area, new_job_state, new_sample, validate_job
from prolab_parser import build_draft_job_from_prolab
from report_builder import MOLD_DESCRIPTIONS, create_report


class ReviewWorkflowTests(unittest.TestCase):
    def test_lab_samples_stay_separate_from_inspection_areas(self):
        job = new_job_state()
        parsed = {
            "metadata": {
                "project_name": "Test Client",
                "test_location": "123 Main St, Dallas, TX 75201",
                "report_date": "September 23, 2026",
            },
            "samples": [
                {
                    "key": "100-1",
                    "location": "OUTDOOR CONTROL",
                    "coc_line": "100-1",
                    "sample_type": "PRO-15",
                    "volume": "75.00 L",
                    "serial_number": "OUT1",
                    "collection_date": "September 22, 2026",
                    "analysis_date": "September 23, 2026",
                    "determination": "CONTROL",
                    "total_spores": 800,
                    "fungi": {"Penicillium/Aspergillus": 210},
                    "is_air": True,
                },
                {
                    "key": "100-2",
                    "location": "INDOORS",
                    "coc_line": "100-2",
                    "sample_type": "PRO-15",
                    "volume": "75.00 L",
                    "serial_number": "IN1",
                    "collection_date": "September 22, 2026",
                    "analysis_date": "September 23, 2026",
                    "determination": "NOT ELEVATED",
                    "total_spores": 643,
                    "fungi": {"Penicillium/Aspergillus": 430},
                    "is_air": True,
                },
            ],
        }

        build_draft_job_from_prolab(job, parsed, MOLD_DESCRIPTIONS.keys())

        self.assertEqual(job["areas"], [])
        indoor = next(s for s in job["samples"] if not s.get("outdoor_control"))
        self.assertEqual(indoor["name"], "INDOORS")
        self.assertIsNone(indoor["area_id"])

    def test_final_validation_requires_rh_and_sample_assignment(self):
        job = new_job_state()
        job["client_name"] = "Client"
        job["address"] = "123 Main"
        job["city"] = "Dallas"
        job["zip"] = "75201"
        job["report_outcome"] = "Mold remediation required"
        job["areas"][0]["name"] = "Laundry"
        job["areas"][0]["finding"] = "Visual mold present"
        job["humidity"] = None

        issues = validate_job(job, lab_pdf_present=True)
        self.assertIn("Indoor RH", issues)

        job["humidity"] = 48
        job["samples"][1]["area_id"] = None
        issues = validate_job(job, lab_pdf_present=True)
        self.assertIn("Assign every indoor/surface sample to an inspection area", issues)

    def test_final_validation_requires_scarlet_style_photos_and_moisture_notes(self):
        job = new_job_state()
        job["client_name"] = "Client"
        job["address"] = "123 Main"
        job["city"] = "Dallas"
        job["state"] = "TX"
        job["zip"] = "75201"
        job["humidity"] = 48
        job["report_outcome"] = "Mold remediation required"
        job["areas"][0]["name"] = "Laundry"
        job["areas"][0]["finding"] = "Visual mold present"
        job["samples"][1]["area_id"] = job["areas"][0]["id"]

        issues = validate_job(job, lab_pdf_present=True)
        self.assertIn("Moisture assessment for Laundry", issues)
        self.assertIn("Property exterior photo", issues)
        self.assertIn("Inspection photos for Laundry", issues)
        self.assertIn("Outdoor control sampling photo", issues)

        job["areas"][0]["moisture_notes"] = "All materials were below 14% moisture content."
        job["photos"] = [
            {"role": "property", "filename": "property.jpg"},
            {"role": "outdoor", "filename": "outdoor.jpg"},
            {"role": "area", "area_id": job["areas"][0]["id"], "filename": "area.jpg"},
        ]
        issues = validate_job(job, lab_pdf_present=True)
        self.assertNotIn("Moisture assessment for Laundry", issues)
        self.assertNotIn("Property exterior photo", issues)
        self.assertNotIn("Inspection photos for Laundry", issues)
        self.assertNotIn("Outdoor control sampling photo", issues)

    def test_report_reflects_current_remediation_outcome_and_area_notes(self):
        area = new_area("Laundry Area")
        area["finding"] = "Active mold growth confirmed"
        area["description"] = "Visible suspect growth was observed behind the washer."
        area["moisture_notes"] = "Elevated moisture was recorded at the wall base."

        outdoor = new_sample("Air Sample", "Outdoor Control", outdoor_control=True)
        outdoor["name"] = "Outdoor Control"
        indoor = new_sample("Air Sample", area_id=area["id"])
        indoor["name"] = "Laundry Air"

        job = {
            "client_name": "Client",
            "address": "123 Main",
            "city": "Dallas",
            "state": "TX",
            "zip": "75201",
            "inspection_date": __import__("datetime").date(2026, 9, 22),
            "report_date": __import__("datetime").date(2026, 9, 23),
            "humidity": 48,
            "temperature": None,
            "general_observations": "",
            "areas": [area],
            "samples": [outdoor, indoor],
            "air_lab_rows": [
                {
                    "id": "a1",
                    "sample_id": outdoor["id"],
                    "fungal_type": "Penicillium/Aspergillus",
                    "spore_count": 210,
                    "interpretation": "Baseline (Reference)",
                },
                {
                    "id": "a2",
                    "sample_id": indoor["id"],
                    "fungal_type": "Penicillium/Aspergillus",
                    "spore_count": 430,
                    "interpretation": "ELEVATED",
                },
            ],
            "surface_lab_rows": [],
            "mold_types": ["Penicillium/Aspergillus"],
            "report_outcome": "Mold remediation required",
        }

        report = create_report(job, {}, None)
        doc = Document(BytesIO(report.getvalue()))
        text = "\n".join(p.text for p in doc.paragraphs)

        self.assertIn("active mold growth was confirmed", text)
        self.assertNotIn("no significant mold contamination was identified", text)
        self.assertIn("Visible suspect growth was observed behind the washer.", text)
        self.assertIn("Elevated moisture was recorded at the wall base.", text)
        self.assertNotIn("Surface Sample Results", text)

        air_table = next(
            table for table in doc.tables
            if table.rows and table.rows[0].cells[0].text == "Fungal Type"
        )
        self.assertEqual(len(air_table.columns), 3)
        self.assertIn("Outdoor Control", air_table.rows[0].cells[1].text)
        self.assertIn("Laundry Air", air_table.rows[0].cells[2].text)


    def test_surface_sample_report_includes_identified_organisms(self):
        area = new_area("Coat Closet")
        area["finding"] = "Visual mold present"

        surface = new_sample("Surface Sample", area_id=area["id"])
        surface["name"] = "COAT CLOSET"
        surface["lab_fungi"] = {
            "Cladosporium": "Present",
            "Curvularia": "Present",
            "Epicoccum": "Present",
            "Hyphae": "Present",
            "Other Ascospores": "Present",
            "Other Basidiospores": "Present",
            "Penicillium/Aspergillus": "Present",
            "Smuts, myxomycetes": "Present",
        }

        job = new_job_state()
        job.update({
            "client_name": "Scarlet Harper",
            "address": "16371 County Road 245",
            "city": "Terrell",
            "state": "TX",
            "zip": "75160",
            "humidity": 50,
            "areas": [area],
            "samples": [surface],
            "air_lab_rows": [],
            "surface_lab_rows": [{
                "id": "surf1",
                "sample_id": surface["id"],
                "result": "UNUSUAL / Mold Present",
            }],
            "mold_types": [
                "Cladosporium",
                "Curvularia",
                "Epicoccum",
                "Hyphae",
                "Other Ascospores",
                "Other Basidiospores",
                "Penicillium/Aspergillus",
                "Smuts, myxomycetes",
            ],
            "report_outcome": "Mold remediation required",
        })

        report = create_report(job, {}, None)
        doc = Document(BytesIO(report.getvalue()))
        surface_table = next(
            table for table in doc.tables
            if table.rows and table.rows[0].cells[0].text == "Sample"
        )

        self.assertEqual(len(surface_table.columns), 4)
        self.assertIn("Cladosporium", surface_table.rows[1].cells[2].text)
        self.assertIn("Penicillium/Aspergillus", surface_table.rows[1].cells[2].text)
        self.assertEqual(surface_table.rows[1].cells[3].text, "UNUSUAL / Mold Present")


if __name__ == "__main__":
    unittest.main()

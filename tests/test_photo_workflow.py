from __future__ import annotations

from datetime import date
from io import BytesIO
import unittest

from docx import Document
from PIL import Image

from models import new_area, new_sample
from photo_service import normalize_report_photo
from report_builder import create_report


def make_image_bytes(width: int = 3200, height: int = 2400) -> bytes:
    image = Image.new("RGB", (width, height), "white")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


class PhotoWorkflowTests(unittest.TestCase):
    def test_normalize_report_photo_converts_and_resizes_phone_image(self):
        normalized = normalize_report_photo(make_image_bytes())

        with Image.open(BytesIO(normalized)) as image:
            self.assertEqual(image.format, "JPEG")
            self.assertLessEqual(max(image.size), 2200)

    def test_report_accepts_multiple_photos_for_one_inspection_area(self):
        area = new_area("Laundry")
        area["finding"] = "No mold detected"
        area["description"] = "Area observation."
        area["moisture_notes"] = "No elevated moisture observed."

        outdoor = new_sample("Air Sample", "Outdoor Control", outdoor_control=True)
        indoor = new_sample("Air Sample", area_id=area["id"])
        indoor["name"] = "Laundry Air"

        job = {
            "client_name": "Photo Test",
            "address": "123 Main",
            "city": "Dallas",
            "state": "TX",
            "zip": "75201",
            "inspection_date": date(2026, 9, 28),
            "report_date": date(2026, 9, 28),
            "humidity": 45,
            "temperature": 74,
            "areas": [area],
            "samples": [outdoor, indoor],
            "air_lab_rows": [],
            "surface_lab_rows": [],
            "mold_types": [],
            "report_outcome": "No significant mold contamination identified",
        }

        image1 = normalize_report_photo(make_image_bytes(800, 600))
        image2 = normalize_report_photo(make_image_bytes(640, 480))

        without_photos = Document(BytesIO(create_report(job, {}, None).getvalue()))
        with_photos = Document(
            BytesIO(
                create_report(
                    job,
                    {
                        area["id"]: [
                            {"content": BytesIO(image1), "caption": "North wall"},
                            {"content": BytesIO(image2), "caption": "Baseboard"},
                        ]
                    },
                    None,
                ).getvalue()
            )
        )

        self.assertEqual(
            len(with_photos.inline_shapes) - len(without_photos.inline_shapes),
            2,
        )
        text = "\n".join(p.text for p in with_photos.paragraphs)
        self.assertIn("North wall", text)
        self.assertIn("Baseboard", text)


if __name__ == "__main__":
    unittest.main()

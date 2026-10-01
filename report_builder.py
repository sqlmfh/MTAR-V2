from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from models import sample_location

BASE_DIR = Path(__file__).resolve().parent
ASSET_DIR = BASE_DIR / "assets"

MOLD_DESCRIPTIONS = {
    "Alternaria": ("A common outdoor mold that often indicates water damage when found indoors. Frequently grows on water-intruded building materials like damp drywall, wood, and textiles. Known allergen and common asthma trigger.", True),
    "Chaetomium": ("A water-indicating mold found on cellulose materials. Should not be observed indoors unless building materials have been wetted.", True),
    "Cladosporium": ("The most common spore type worldwide. Commonly found on wood and wallboard. Known allergen but also common outdoors.", False),
    "Curvularia": ("A common outdoor mold and plant pathogen. When found indoors, it can grow on various building materials and may indicate moisture issues. Known allergen.", False),
    "Epicoccum": ("A widespread outdoor fungus commonly found on plant debris. When found indoors, it is typically associated with water-damaged building materials like drywall or paper. Known allergen.", False),
    "Hyphae": ("Fragments of fungal structures (the root system of mold). Finding these in high concentrations indoors strongly indicates active mold growth and amplification.", True),
    "Other Ascospores": ("Spores from a large group of fungi common everywhere outdoors. When found indoors in higher concentrations, they may indicate moisture issues.", False),
    "Other Basidiospores": ("A common outdoor spore type originating from mushrooms and bracket fungi. High indoor concentrations can sometimes indicate wood decay or structural moisture problems.", False),
    "Penicillium/Aspergillus": ("The most common mold species in indoor air samples. Often associated with water damage and elevated humidity. Known allergen (Type I and Type III).", False),
    "Smuts, myxomycetes": ("Common spores found outdoors on plants, grasses, and in soil. They rarely grow indoors; their presence is usually due to normal infiltration of outdoor air.", False),
    "Stachybotrys": ("Known as 'black mold.' Requires high water content to grow. A water-indicating mold that produces mycotoxins. Professional remediation required.", True),
    "Trichocladium": ("Rarely seen in the air, this mold grows most commonly on decaying or wetted wood.", False),
}


def set_cell_shading(cell, color: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def make_tight(para):
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.15
    return para


def make_top_tight(para):
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.line_spacing = 1.15
    return para


def add_floating_image(paragraph, image_path: Path, width, x_pos: float, y_pos: float):
    if not image_path.exists():
        return
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=width)
    inline = shape._inline
    extent = inline.extent
    graphic_xml = inline.graphic.xml
    x_emu = int(x_pos * 914400)
    y_emu = int(y_pos * 914400)
    anchor_xml = (
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1" '
        f'{nsdecls("wp", "a", "pic", "r")}>'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x_emu}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y_emu}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{extent.cx}" cy="{extent.cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="1" name="FixedImage"/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'{graphic_xml}'
        f'</wp:anchor>'
    )
    anchor = parse_xml(anchor_xml)
    inline.getparent().replace(inline, anchor)


def add_page_number(run):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def _sample_map(job: dict) -> dict[str, dict]:
    return {s["id"]: s for s in job.get("samples", [])}


def _sample_display_name(sample: dict) -> str:
    name = (sample.get("name") or "").strip()
    if name:
        return name
    if sample.get("outdoor_control"):
        return "Outdoor Control"
    return sample.get("lab_location") or "Unnamed Sample"


def _assigned_area_name(sample: dict, areas: list[dict]) -> str:
    if sample.get("outdoor_control"):
        return "Outdoor Control"
    area_id = sample.get("area_id")
    for area in areas:
        if area.get("id") == area_id:
            return area.get("name") or "Unnamed Area"
    return "Unassigned"


def _lab_rows_for_report(job: dict) -> list[dict]:
    samples = _sample_map(job)
    rows = []
    for row in job.get("air_lab_rows", []):
        sample = samples.get(row.get("sample_id"), {})
        rows.append(
            {
                "location": sample_location(sample, job.get("areas", [])),
                "fungal_type": row.get("fungal_type", ""),
                "spore_count": row.get("spore_count", 0),
                "interpretation": row.get("interpretation", ""),
            }
        )
    return rows


def _surface_rows_for_report(job: dict) -> list[dict]:
    samples = _sample_map(job)
    rows = []
    for row in job.get("surface_lab_rows", []):
        sample = samples.get(row.get("sample_id"), {})
        rows.append(
            {
                "location": sample_location(sample, job.get("areas", [])),
                "result": row.get("result", ""),
            }
        )
    return rows


def _photo_entries(value) -> list[dict]:
    """Normalize legacy single-photo inputs and new multi-photo entries."""
    if not value:
        return []
    if isinstance(value, (list, tuple)):
        entries = []
        for item in value:
            if isinstance(item, dict):
                entries.append(item)
            else:
                entries.append({"content": item, "caption": ""})
        return entries
    if isinstance(value, dict) and "content" in value:
        return [value]
    return [{"content": value, "caption": ""}]


def _add_report_photo(doc: Document, entry: dict, width=Inches(3.4)) -> None:
    content = entry.get("content")
    if not content:
        return
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(content, width=width)
        caption = (entry.get("caption") or "").strip()
        if caption:
            cp = make_tight(doc.add_paragraph())
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cp.add_run(caption)
            run.italic = True
            run.font.size = Pt(9)
    except Exception:
        pass


def create_report(job: dict, photos: dict, lab_pdf_bytes: bytes | None = None) -> BytesIO:
    """Generate the V2 draft Word report from one structured job record.

    Photos supports both the legacy single-file shape and the new list-of-photo
    entries used by the persistent automation workflow.
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    section = doc.sections[0]
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())
    add_floating_image(
        footer_para,
        ASSET_DIR / "Azeem_TDLR_Signature.png",
        width=Inches(0.69),
        x_pos=7.7,
        y_pos=0.1,
    )

    contact = make_tight(doc.add_paragraph())
    add_floating_image(contact, ASSET_DIR / "MTAR_logo.png", width=Inches(2.92), x_pos=1.0, y_pos=0.75)
    contact.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    contact.add_run("Mold Testing and Removal\n")
    contact.add_run("2031 John West Rd. #119\n")
    contact.add_run("Dallas, TX 75228\n")
    contact.add_run("(817) 718-5086\n")
    contact.add_run("help@moldtestingandremoval.com")

    make_tight(doc.add_paragraph())
    title = make_tight(doc.add_paragraph())
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MOLD ASSESSMENT REPORT")
    run.font.name = "Bebas Neue"
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(24, 64, 88)
    make_tight(doc.add_paragraph())

    property_photos = _photo_entries(photos.get("property"))
    if property_photos:
        cover_entry = property_photos[0]
        try:
            p = make_tight(doc.add_paragraph())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(cover_entry.get("content"), width=Inches(5))
            cover_caption = (cover_entry.get("caption") or "").strip()
            if cover_caption:
                cp = make_tight(doc.add_paragraph())
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cp.add_run(cover_caption)
                rr.italic = True
                rr.font.size = Pt(9)
        except Exception:
            pass

    make_tight(doc.add_paragraph())
    info = make_tight(doc.add_paragraph())
    run = info.add_run("Client & Property:\n")
    run.font.name = "Bebas Neue"
    run.bold = True
    run.font.color.rgb = RGBColor(24, 64, 88)
    run.font.size = Pt(15)
    info.add_run(f"{job['client_name']}\n{job['address']}\n{job['city']}, {job['state']} {job['zip']}\n")

    info2 = make_tight(doc.add_paragraph())
    run = info2.add_run("Assessment Date: ")
    run.font.name = "Bebas Neue"
    run.bold = True
    run.font.color.rgb = RGBColor(24, 64, 88)
    run.font.size = Pt(15)
    info2.add_run(job["inspection_date"].strftime("%B %d, %Y"))

    info3 = make_tight(doc.add_paragraph())
    run = info3.add_run("Report Date: ")
    run.font.name = "Bebas Neue"
    run.bold = True
    run.font.color.rgb = RGBColor(24, 64, 88)
    run.font.size = Pt(15)
    info3.add_run(job["report_date"].strftime("%B %d, %Y"))

    info4 = make_tight(doc.add_paragraph())
    run = info4.add_run("Samples Taken:\n")
    run.font.name = "Bebas Neue"
    run.bold = True
    run.font.color.rgb = RGBColor(24, 64, 88)
    run.font.size = Pt(15)
    for index, sample in enumerate(job.get("samples", []), 1):
        prefix = "Exterior control sample" if sample.get("outdoor_control") else f"Sample {index - 1}"
        sample_name = _sample_display_name(sample)
        area_name = _assigned_area_name(sample, job.get("areas", []))
        if sample.get("outdoor_control"):
            info4.add_run(f"{prefix}: {sample_name} ({sample['type']})\n")
        else:
            info4.add_run(f"{prefix}: {sample_name} ({sample['type']}) — assigned to {area_name}\n")

    doc.add_page_break()
    letter_header = make_tight(doc.add_paragraph())
    r = letter_header.add_run("State Licensed Mold Assessment Consultant:\n")
    r.bold = True
    letter_header.add_run("Azeem Iqbal — TDLR MAC #2189\n\n")
    r = letter_header.add_run("Report Date:\n")
    r.bold = True
    letter_header.add_run(job["report_date"].strftime("%B %d, %Y"))
    make_tight(doc.add_paragraph())
    doc.add_paragraph("To whom it may concern,")

    doc.add_paragraph(
        f"Mold Testing and Removal was hired to conduct a mold assessment at the property located at "
        f"{job['address']}, {job['city']}, {job['state']} {job['zip']}. The purpose of this assessment was to "
        "evaluate the indoor air quality, identify potential sources of fungal growth, and provide recommendations for remediation."
    )
    has_air_samples = any(s.get("type") == "Air Sample" for s in job.get("samples", []))
    has_surface_samples = bool(job.get("surface_lab_rows"))
    if has_air_samples and has_surface_samples:
        sampling_text = "bioaerosol (air) and surface samples"
    elif has_surface_samples:
        sampling_text = "surface samples"
    else:
        sampling_text = "bioaerosol (air) samples"

    doc.add_paragraph(
        "The assessment included a visual inspection, moisture mapping using a Protimeter Moisture Meter, "
        f"and the collection of {sampling_text}. Samples were collected from the assessed areas"
        + (" and the exterior for control purposes." if has_air_samples else ".")
    )
    doc.add_paragraph("The samples were sent to PRO-LAB, an accredited laboratory, for mold/fungi analysis.")

    report_outcome = job.get("report_outcome", "Pending consultant review")
    remediation_required = report_outcome == "Mold remediation required"
    pending_review = report_outcome in ("", "Pending consultant review")

    if remediation_required:
        results_p = doc.add_paragraph()
        results_p.add_run("Based on the laboratory results and visual inspection, ")
        r = results_p.add_run("active mold growth was confirmed")
        r.bold = True
        r.italic = True
        results_p.add_run(" in the following areas:")
        for area in job.get("areas", []):
            if not area.get("name"):
                continue
            bullet = doc.add_paragraph(style="List Bullet")
            bullet.add_run(f"{area['name']} — {area.get('finding', '')}")
        notification = doc.add_paragraph()
        r = notification.add_run(
            "This letter serves as official notification that professional mold remediation is required to return "
            "the property to a normal fungal ecology (Condition 1). The property should be remediated by a State "
            "Licensed Mold Remediation Contractor (MRC) in accordance with the Texas Mold Assessment and Remediation Rules (TMARR)."
        )
        r.bold = True
    elif pending_review:
        results_p = doc.add_paragraph()
        r = results_p.add_run("DRAFT - CONSULTANT REVIEW REQUIRED")
        r.bold = True
        r.italic = True
        r.font.color.rgb = RGBColor(220, 53, 69)
        results_p.add_run(
            ". Laboratory data has been imported, but visual findings, moisture observations, "
            "area findings, and the final professional conclusion must be reviewed before this report is finalized."
        )
    else:
        results_p = doc.add_paragraph()
        results_p.add_run("Based on the laboratory results and visual inspection, ")
        r = results_p.add_run("no significant mold contamination was identified")
        r.bold = True
        r.italic = True
        results_p.add_run(". The indoor spore counts are within normal parameters compared to the outdoor control sample.")

    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    sig = make_tight(doc.add_paragraph())
    r = sig.add_run("Azeem Iqbal")
    r.bold = True
    r.font.size = Pt(12)
    sig_para = make_tight(doc.add_paragraph())
    if (ASSET_DIR / "Signature.png").exists():
        sig_para.add_run().add_picture(str(ASSET_DIR / "Signature.png"), width=Inches(0.9))
    make_tight(doc.add_paragraph("State of Texas Licensed Mold Assessment Consultant"))
    make_tight(doc.add_paragraph("TDLR MAC #2189 (Exp. 10/24/2027)"))

    doc.add_page_break()
    obs_title = make_tight(doc.add_paragraph())
    r = obs_title.add_run("Visual Observations & Moisture Readings")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)

    env_p = make_top_tight(doc.add_paragraph())
    r = env_p.add_run("Environmental Conditions: ")
    r.bold = True
    humidity = job.get("humidity")
    if humidity is None:
        r = env_p.add_run("Indoor relative humidity was not entered in this draft. Consultant review required.")
        r.italic = True
        r.font.color.rgb = RGBColor(220, 53, 69)
    else:
        env_p.add_run("The indoor relative humidity (rH) was recorded at ")
        r = env_p.add_run(f"{humidity}%")
        r.bold = True
        env_p.add_run(", which is ")
        if humidity > 50:
            r.font.color.rgb = RGBColor(220, 53, 69)
            env_p.add_run("above the recommended range (30-50%) and conducive to microbial growth.")
        else:
            env_p.add_run("within the recommended range (30-50%).")

    ocs_title = make_tight(doc.add_paragraph())
    r = ocs_title.add_run("Outdoor Control Sample")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)
    make_top_tight(doc.add_paragraph("An air sample is taken outside to serve as a baseline for all other air samples to be compared against."))

    for area in job.get("areas", []):
        if not area.get("name"):
            continue
        doc.add_page_break()
        area_title = make_tight(doc.add_paragraph())
        r = area_title.add_run(area["name"])
        r.font.name = "Bebas Neue"
        r.bold = True
        r.font.color.rgb = RGBColor(24, 64, 88)
        r.font.size = Pt(17)
        if area.get("lab_summary"):
            doc.add_paragraph(area["lab_summary"])
        if area.get("description"):
            p = doc.add_paragraph()
            rr = p.add_run("Visual Observations: ")
            rr.bold = True
            p.add_run(area["description"])
        if area.get("moisture_notes"):
            p = doc.add_paragraph()
            rr = p.add_run("Moisture Assessment: ")
            rr.bold = True
            p.add_run(area["moisture_notes"])
        area_photos = _photo_entries(photos.get(area["id"]))
        for entry in area_photos:
            _add_report_photo(doc, entry)

    doc.add_page_break()
    lab_title = make_tight(doc.add_paragraph())
    r = lab_title.add_run("Laboratory Results Analysis")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)
    make_top_tight(doc.add_paragraph(
        "Samples were submitted to PRO-LAB (an accredited laboratory) for analysis. "
        "The following summarizes the findings compared to the outdoor control sample."
    ))

    air_rows = job.get("air_lab_rows", [])
    air_samples = [
        sample for sample in job.get("samples", [])
        if sample.get("type") == "Air Sample"
    ]
    if air_rows and air_samples:
        air_title = make_tight(doc.add_paragraph())
        r = air_title.add_run("Air Sample Comparison (Bioaerosol)")
        r.font.name = "Bebas Neue"
        r.bold = True
        r.font.color.rgb = RGBColor(24, 64, 88)
        r.font.size = Pt(17)

        # PRO-LAB-style comparison: fungal types down the left, samples side by side.
        air_table = doc.add_table(rows=1, cols=1 + len(air_samples))
        air_table.style = "Table Grid"
        header = air_table.rows[0].cells
        header[0].text = "Fungal Type"
        header[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(header[0], "D5E8F0")

        for index, sample in enumerate(air_samples, 1):
            sample_name = _sample_display_name(sample)
            area_name = _assigned_area_name(sample, job.get("areas", []))
            header_text = sample_name
            if not sample.get("outdoor_control"):
                header_text += f"\n{area_name}"
            header[index].text = header_text
            header[index].paragraphs[0].runs[0].bold = True
            set_cell_shading(header[index], "D5E8F0")

        species = []
        lookup = {}
        for lab_row in air_rows:
            fungus = lab_row.get("fungal_type", "")
            if fungus and fungus not in species:
                species.append(fungus)
            lookup[(fungus, lab_row.get("sample_id"))] = lab_row

        for fungus in species:
            row = air_table.add_row()
            row.cells[0].text = fungus
            for index, sample in enumerate(air_samples, 1):
                lab_row = lookup.get((fungus, sample.get("id")))
                if not lab_row:
                    row.cells[index].text = "—"
                    continue
                count = lab_row.get("spore_count", 0)
                interpretation = lab_row.get("interpretation", "")
                row.cells[index].text = f"{count}\n{interpretation}"
                if interpretation.upper() == "ELEVATED":
                    set_cell_shading(row.cells[index], "FFCCCC")
                    for run in row.cells[index].paragraphs[0].runs:
                        run.font.color.rgb = RGBColor(220, 53, 69)
                        run.bold = True

        if any(sample.get("lab_total_spores") is not None for sample in air_samples):
            total = air_table.add_row()
            total.cells[0].text = "TOTAL SPORES"
            total.cells[0].paragraphs[0].runs[0].bold = True
            for index, sample in enumerate(air_samples, 1):
                value = sample.get("lab_total_spores")
                total.cells[index].text = "—" if value is None else str(value)
                total.cells[index].paragraphs[0].runs[0].bold = True

    surface_rows = _surface_rows_for_report(job)
    if surface_rows:
        doc.add_paragraph()
        surface_title = make_tight(doc.add_paragraph())
        r = surface_title.add_run("Surface Sample Results")
        r.font.name = "Bebas Neue"
        r.bold = True
        r.font.color.rgb = RGBColor(24, 64, 88)
        r.font.size = Pt(17)
        surface_table = doc.add_table(rows=1, cols=4)
        surface_table.style = "Table Grid"
        for i, text in enumerate(["Sample", "Assigned Area", "Mold Identified", "Result"]):
            cell = surface_table.rows[0].cells[i]
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
            set_cell_shading(cell, "D5E8F0")

        samples = _sample_map(job)
        for source_row in job.get("surface_lab_rows", []):
            sample = samples.get(source_row.get("sample_id"), {})
            row = surface_table.add_row()
            row.cells[0].text = _sample_display_name(sample)
            row.cells[1].text = _assigned_area_name(sample, job.get("areas", []))
            surface_fungi = [
                name
                for name, value in sample.get("lab_fungi", {}).items()
                if value not in (None, "", 0, False)
            ]
            row.cells[2].text = ", ".join(surface_fungi) if surface_fungi else "—"
            row.cells[3].text = source_row.get("result", "")
            if "UNUSUAL" in row.cells[3].text.upper() or "MOLD PRESENT" in row.cells[3].text.upper():
                set_cell_shading(row.cells[2], "FFCCCC")
                set_cell_shading(row.cells[3], "FFCCCC")
                row.cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(220, 53, 69)
                row.cells[3].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    mold_title = make_tight(doc.add_paragraph())
    r = mold_title.add_run("Mold Types Identified")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)
    for mold_type in job.get("mold_types", []):
        description, dangerous = MOLD_DESCRIPTIONS[mold_type]
        p = doc.add_paragraph()
        rr = p.add_run(f"{mold_type}: ")
        rr.bold = True
        if dangerous:
            rr.font.color.rgb = RGBColor(220, 53, 69)
        p.add_run(description)

    doc.add_page_break()
    conc_title = make_tight(doc.add_paragraph())
    r = conc_title.add_run("Conclusions")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)
    doc.add_paragraph("Based on the visual inspection, moisture readings, and laboratory results, the following conclusions are made:")
    for i, area in enumerate([a for a in job.get("areas", []) if a.get("name")], 1):
        p = doc.add_paragraph()
        rr = p.add_run(f"{i}. {area['name']}: ")
        rr.bold = True
        p.add_run(area.get("finding", ""))

    doc.add_paragraph()
    rec_title = make_tight(doc.add_paragraph())
    r = rec_title.add_run("Recommendations")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)

    if remediation_required:
        doc.add_paragraph("To return the property to a normal fungal ecology (Condition 1), the following remediation steps are recommended:")
        recommendations = [
            ("Professional Remediation", "Hire a State Licensed Mold Remediation Contractor (MRC) to prepare a work plan based on a Mold Remediation Protocol prepared by a TDLR Mold Assessment Consultant."),
            ("Containment", "Establish critical barriers (polyethylene sheeting) around affected areas to prevent spore dispersion. Establish negative air pressure."),
            ("Removal", "Remove and discard affected drywall and materials. Continue removal 2 feet beyond visible growth."),
            ("Cleaning", "HEPA vacuum and damp-wipe all remaining structural surfaces within the containment."),
            ("Humidity Control", "Dehumidification is required to lower the indoor RH to between 30-50%."),
            ("Clearance Testing", "After remediation, a Post-Remediation Assessment (clearance test) must be performed by a TDLR Mold Assessment Consultant."),
        ]
    elif pending_review:
        doc.add_paragraph(
            "Recommendations are intentionally withheld in this draft until the licensed consultant completes review."
        )
        recommendations = []
    else:
        doc.add_paragraph("Based on the findings, the following recommendations are made:")
        recommendations = [
            ("Humidity Control", "Maintain indoor relative humidity between 30-50% to prevent future mold growth."),
            ("Regular Inspection", "Periodically check areas prone to moisture for signs of water intrusion or condensation."),
            ("Ventilation", "Ensure proper ventilation in bathrooms, kitchens, and laundry areas."),
        ]
    for title_text, body in recommendations:
        p = doc.add_paragraph(style="List Bullet")
        rr = p.add_run(f"{title_text}: ")
        rr.bold = True
        p.add_run(body)

    doc.add_paragraph()
    compliance = doc.add_paragraph()
    r = compliance.add_run(
        "This report is generated in accordance with the Texas Mold Assessment and Remediation Rules (TMARR). "
        "Limitations: This inspection is limited to the areas accessible at the time of inspection."
    )
    r.italic = True
    r.font.size = Pt(9)

    doc.add_page_break()
    terms_title = make_tight(doc.add_paragraph())
    r = terms_title.add_run("Terms and Conditions")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(17)
    terms = [
        ("Inspection Limitation", "This inspection and the information set forth in the report is provided solely for the purpose of verifying that certain structural or physical characteristics exist at the Location Address listed. The undersigned and company representative does not make a health or safety certification or warranty, express or implied, of any kind."),
        ("Limitation of Liability", "The Client agrees that Inspector's liability for errors and/or omissions shall be limited to the maximum of a full refund of the fee paid for the inspection. The Client agrees to assume all risk of loss which exceeds the fee paid."),
        ("Sampling Limitations", "Mold spore sampling results represent conditions at the time and location of sampling only. Conditions can change rapidly due to environmental factors, occupant activities, and remediation efforts."),
        ("Health Disclaimer", "This report does not constitute medical advice. Individuals with health concerns related to potential mold exposure should consult with a qualified healthcare professional."),
        ("Report Usage", "This report is prepared exclusively for the named client and may not be reproduced or distributed to third parties without written consent from Mold Testing and Removal."),
    ]
    for title_text, body in terms:
        p = doc.add_paragraph()
        rr = p.add_run(f"{title_text}: ")
        rr.bold = True
        rr.font.size = Pt(10)
        p.add_run(body).font.size = Pt(10)

    doc.add_paragraph()
    lab_ref = doc.add_paragraph()
    lab_ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = lab_ref.add_run("— Laboratory Report Attached —")
    r.bold = True
    r.italic = True
    lab_ref2 = doc.add_paragraph()
    lab_ref2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lab_ref2.add_run("PRO-LAB Certificate of Mold Analysis follows this page")

    doc.add_paragraph()
    company_footer = doc.add_paragraph()
    company_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = company_footer.add_run("Mold Testing and Removal\n")
    r.font.name = "Bebas Neue"
    r.bold = True
    r.font.color.rgb = RGBColor(24, 64, 88)
    r.font.size = Pt(20)
    company_footer.add_run("2031 John West Rd. #119 | Dallas, TX 75228\n")
    company_footer.add_run("(817) 718-5086 | help@moldtestingandremoval.com")

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
